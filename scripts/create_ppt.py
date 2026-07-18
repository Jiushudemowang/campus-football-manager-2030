#!/usr/bin/env python3
"""绿茵纪比赛答辩 PPT 生成脚本"""

from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN
from pptx.enum.shapes import MSO_SHAPE
import os

# 颜色方案：绿茵主题
BG_DARK   = RGBColor(0x0B, 0x14, 0x0B)   # 深绿茵
BG_CARD   = RGBColor(0x14, 0x24, 0x16)   # 卡片色
GREEN     = RGBColor(0x22, 0xC5, 0x5E)   # 草坪绿
YELLOW    = RGBColor(0xFB, 0xBF, 0x24)   # 暖沙金
ORANGE    = RGBColor(0xF9, 0x73, 0x16)   # 夕阳橙
WHITE     = RGBColor(0xF5, 0xF5, 0xF0)
GRAY      = RGBColor(0x94, 0xA3, 0xB8)

prs = Presentation()
prs.slide_width  = Inches(13.333)
prs.slide_height = Inches(7.5)

def add_bg(slide, color=BG_DARK):
    bg = slide.background
    fill = bg.fill
    fill.solid()
    fill.fore_color.rgb = color

def add_rect(slide, left, top, width, height, color, radius=None):
    shape = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE if radius else MSO_SHAPE.RECTANGLE, left, top, width, height)
    shape.fill.solid()
    shape.fill.fore_color.rgb = color
    shape.line.fill.background()
    return shape

def add_textbox(slide, left, top, width, height, text, font_size=18, color=WHITE, bold=False, align=PP_ALIGN.LEFT, font_name='Microsoft YaHei'):
    txBox = slide.shapes.add_textbox(left, top, width, height)
    tf = txBox.text_frame
    tf.word_wrap = True
    p = tf.paragraphs[0]
    p.text = text
    p.font.size = Pt(font_size)
    p.font.color.rgb = color
    p.font.bold = bold
    p.font.name = font_name
    p.alignment = align
    return txBox

def add_accent_line(slide, left, top, width, color=GREEN):
    shape = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, left, top, width, Pt(3))
    shape.fill.solid()
    shape.fill.fore_color.rgb = color
    shape.line.fill.background()
    return shape

# ═══════════════════════════════════
# Slide 1: 封面
# ═══════════════════════════════════
slide = prs.slides.add_slide(prs.slide_layouts[6])
add_bg(slide)
add_rect(slide, Inches(0), Inches(0), Inches(13.333), Inches(0.08), GREEN)
add_rect(slide, Inches(0), Inches(7.42), Inches(13.333), Inches(0.08), GREEN)
add_rect(slide, Inches(0), Inches(0), Inches(0.12), Inches(7.5), YELLOW)

add_textbox(slide, Inches(1.5), Inches(1.5), Inches(10), Inches(1.2),
            "绿茵纪", font_size=72, color=WHITE, bold=True)
add_textbox(slide, Inches(1.5), Inches(2.8), Inches(10), Inches(0.6),
            "AI 体育叙事引擎", font_size=32, color=GREEN, bold=True)
add_accent_line(slide, Inches(1.5), Inches(3.6), Inches(2.5))
add_textbox(slide, Inches(1.5), Inches(4.0), Inches(10), Inches(0.4),
            "第二届 AIGC 应用大赛  ·  创作赛道  ·  AI 赋能体育传播", font_size=16, color=GRAY)
add_textbox(slide, Inches(1.5), Inches(5.8), Inches(8), Inches(0.4),
            "用 AI 记录、重构与传播民间体育历史", font_size=14, color=GRAY)
add_textbox(slide, Inches(1.5), Inches(6.2), Inches(5), Inches(0.4),
            "2026年6月", font_size=12, color=GRAY)

# 右侧装饰圆
circle = slide.shapes.add_shape(MSO_SHAPE.OVAL, Inches(9.0), Inches(1.0), Inches(4.5), Inches(4.5))
circle.fill.solid()
circle.fill.fore_color.rgb = GREEN
circle.line.fill.background()
circle2 = slide.shapes.add_shape(MSO_SHAPE.OVAL, Inches(9.8), Inches(2.2), Inches(3.0), Inches(3.0))
circle2.fill.solid()
circle2.fill.fore_color.rgb = BG_DARK
circle2.line.fill.background()

# ═══════════════════════════════════
# Slide 2: 目录
# ═══════════════════════════════════
slide = prs.slides.add_slide(prs.slide_layouts[6])
add_bg(slide)
add_accent_line(slide, Inches(0.8), Inches(0.7), Inches(0.6))
add_textbox(slide, Inches(0.8), Inches(0.8), Inches(10), Inches(0.5),
            "CONTENTS / 目录", font_size=14, color=GREEN, bold=True)

toc_items = [
    ("01", "项目背景", "体育传播的痛点与机遇"),
    ("02", "解决方案", "AI 互动叙事引擎"),
    ("03", "技术架构", "多模态 AIGC 融合"),
    ("04", "作品展示", "核心功能与数据可视化"),
    ("05", "创新价值", "社会意义与未来规划"),
]

for i, (num, title, desc) in enumerate(toc_items):
    y = Inches(1.7) + Inches(1.05) * i
    add_rect(slide, Inches(1.2), y, Inches(10.8), Inches(0.9), BG_CARD, radius=True)
    add_textbox(slide, Inches(1.6), y + Inches(0.1), Inches(0.8), Inches(0.5),
                num, font_size=32, color=GREEN, bold=True)
    add_textbox(slide, Inches(2.5), y + Inches(0.10), Inches(4), Inches(0.35),
                title, font_size=20, color=WHITE, bold=True)
    add_textbox(slide, Inches(2.5), y + Inches(0.48), Inches(7), Inches(0.3),
                desc, font_size=12, color=GRAY)

# ═══════════════════════════════════
# Slide 3: 项目背景
# ═══════════════════════════════════
slide = prs.slides.add_slide(prs.slide_layouts[6])
add_bg(slide)
add_accent_line(slide, Inches(0.8), Inches(0.7), Inches(0.6))
add_textbox(slide, Inches(0.8), Inches(0.8), Inches(10), Inches(0.5),
            "01 / 项目背景", font_size=14, color=GREEN, bold=True)
add_textbox(slide, Inches(0.8), Inches(1.1), Inches(11), Inches(0.8),
            "体育传播的「遗忘危机」", font_size=28, color=WHITE, bold=True)

cards_data = [
    ("⚽ 职业体育", "头部赛事占据绝大部分传播资源", ["奥运/世界杯/中超等巨头赛道", "技术方案同质化严重", "中小赛事难以被看见"]),
    ("🏫 民间体育", "大量真实故事缺少记录与传播工具", ["校园球队、社区赛事、业余俱乐部", "历史依赖个人记忆和零散影像", "缺乏情感化、沉浸化表达"]),
    ("🤖 AIGC 机遇", "AI 让历史叙事变得可参与、可传播", ["多模态生成降低制作成本", "互动叙事提升情感共鸣", "Web 应用便于分享传播"]),
]

for i, (title, subtitle, items) in enumerate(cards_data):
    x = Inches(0.8) + Inches(4.0) * i
    y = Inches(2.1)
    add_rect(slide, x, y, Inches(3.7), Inches(4.6), BG_CARD, radius=True)
    colors = [YELLOW, ORANGE, GREEN]
    add_rect(slide, x, y, Inches(3.7), Pt(4), colors[i])
    add_textbox(slide, x + Inches(0.3), y + Inches(0.3), Inches(3.1), Inches(0.4),
                title, font_size=18, color=WHITE, bold=True)
    add_textbox(slide, x + Inches(0.3), y + Inches(0.75), Inches(3.1), Inches(0.5),
                subtitle, font_size=13, color=colors[i])
    for j, item in enumerate(items):
        iy = y + Inches(1.45) + Inches(0.65) * j
        add_textbox(slide, x + Inches(0.3), iy, Inches(0.12), Inches(0.3),
                    "▸", font_size=10, color=colors[i])
        add_textbox(slide, x + Inches(0.55), iy, Inches(3.0), Inches(0.5),
                    item, font_size=13, color=GRAY)

# ═══════════════════════════════════
# Slide 4: 解决方案
# ═══════════════════════════════════
slide = prs.slides.add_slide(prs.slide_layouts[6])
add_bg(slide)
add_accent_line(slide, Inches(0.8), Inches(0.7), Inches(0.6))
add_textbox(slide, Inches(0.8), Inches(0.8), Inches(10), Inches(0.5),
            "02 / 解决方案", font_size=14, color=GREEN, bold=True)
add_textbox(slide, Inches(0.8), Inches(1.1), Inches(11), Inches(0.8),
            "绿茵纪：AI 体育叙事引擎", font_size=28, color=WHITE, bold=True)

add_textbox(slide, Inches(1.0), Inches(2.0), Inches(11.3), Inches(0.8),
            "一个基于 AIGC 的互动体育历史叙事平台。以华中科技大学新闻男足 2016–2026 十年队史为首个案例，"
            "让用户选择不同球员视角，亲历一支草根球队从诞生、低谷、崛起到传承的完整历程。",
            font_size=16, color=GRAY, align=PP_ALIGN.CENTER)

features = [
    ("🎮", "互动叙事", "多视角、多分支、多结局"),
    ("🎨", "AI 视觉", "像素风格场景与角色生成"),
    ("📊", "数据可视化", "十年战绩时间线"),
    ("🌐", "Web 应用", "跨平台、易分享"),
    ("📚", "历史档案", "真实故事数字化活化"),
]

for i, (icon, title, desc) in enumerate(features):
    x = Inches(0.6) + Inches(2.5) * i
    y = Inches(3.2)
    add_rect(slide, x, y, Inches(2.2), Inches(2.6), BG_CARD, radius=True)
    add_rect(slide, x, y, Inches(2.2), Pt(3), GREEN)
    add_textbox(slide, x, y + Inches(0.2), Inches(2.2), Inches(0.6),
                icon, font_size=36, color=WHITE, bold=True, align=PP_ALIGN.CENTER)
    add_textbox(slide, x + Inches(0.15), y + Inches(0.9), Inches(1.9), Inches(0.4),
                title, font_size=14, color=GREEN, bold=True, align=PP_ALIGN.CENTER)
    add_textbox(slide, x + Inches(0.15), y + Inches(1.35), Inches(1.9), Inches(0.8),
                desc, font_size=11, color=GRAY, align=PP_ALIGN.CENTER)

# ═══════════════════════════════════
# Slide 5: 技术架构
# ═══════════════════════════════════
slide = prs.slides.add_slide(prs.slide_layouts[6])
add_bg(slide)
add_accent_line(slide, Inches(0.8), Inches(0.7), Inches(0.6))
add_textbox(slide, Inches(0.8), Inches(0.8), Inches(10), Inches(0.5),
            "03 / 技术架构", font_size=14, color=GREEN, bold=True)
add_textbox(slide, Inches(0.8), Inches(1.1), Inches(11), Inches(0.8),
            "多模态 AIGC 深度融合", font_size=28, color=WHITE, bold=True)

layers = [
    ("表现层", "React + TypeScript + Vite", "互动叙事界面、角色选择、对话系统、结局展示", YELLOW),
    ("AI 生成层", "Trae API / 即梦", "像素风格场景图、角色头像、背景图自动生成", GREEN),
    ("数据层", "Zustand + JSON", "角色档案、故事线、分支选择、结局判定、存档系统", ORANGE),
    ("可视化层", "ECharts", "十年战绩曲线、关键事件标注、历史节点卡片", YELLOW),
]

for i, (name, tech, func, color) in enumerate(layers):
    y = Inches(2.0) + Inches(1.2) * i
    add_rect(slide, Inches(1.5), y, Inches(10.3), Inches(1.0), BG_CARD, radius=True)
    add_rect(slide, Inches(1.5), y, Pt(5), Inches(1.0), color)
    add_textbox(slide, Inches(1.8), y + Inches(0.1), Inches(1.5), Inches(0.3),
                name, font_size=16, color=color, bold=True)
    add_textbox(slide, Inches(3.3), y + Inches(0.1), Inches(3.0), Inches(0.3),
                tech, font_size=13, color=WHITE)
    add_textbox(slide, Inches(3.3), y + Inches(0.55), Inches(8.5), Inches(0.35),
                func, font_size=11, color=GRAY)

for i in range(3):
    y = Inches(2.6) + Inches(1.2) * i
    add_textbox(slide, Inches(12.0), y, Inches(0.5), Inches(0.5),
                "↓", font_size=18, color=GREEN, bold=True, align=PP_ALIGN.CENTER)

# ═══════════════════════════════════
# Slide 6: 作品展示
# ═══════════════════════════════════
slide = prs.slides.add_slide(prs.slide_layouts[6])
add_bg(slide)
add_accent_line(slide, Inches(0.8), Inches(0.7), Inches(0.6))
add_textbox(slide, Inches(0.8), Inches(0.8), Inches(10), Inches(0.5),
            "04 / 作品展示", font_size=14, color=GREEN, bold=True)
add_textbox(slide, Inches(0.8), Inches(1.1), Inches(11), Inches(0.8),
            "核心功能", font_size=28, color=WHITE, bold=True)

showcases = [
    ("6 位核心角色", "吴志鸣、肖潇、吾尔肯、赵凌冬、杨云帆、王楷硕", "每位代表一个时代"),
    ("多分支叙事", "每个角色 4–5 个场景，含关键选择", "选择影响结局走向"),
    ("3 种结局", "光辉 / 平凡 / 遗憾", "基于选择权重动态判定"),
    ("存档系统", "手动存档 + 自动保存", "支持随时继续体验"),
    ("十年战绩", "ECharts 数据可视化", "完整呈现球队起伏历程"),
]

for i, (title, desc, detail) in enumerate(showcases):
    x = Inches(0.8) + Inches(6.0) * (i % 2)
    y = Inches(2.0) + Inches(1.35) * (i // 2)
    add_rect(slide, x, y, Inches(5.7), Inches(1.15), BG_CARD, radius=True)
    add_rect(slide, x, y, Inches(5.7), Pt(3), GREEN)
    add_textbox(slide, x + Inches(0.25), y + Inches(0.15), Inches(5.2), Inches(0.35),
                title, font_size=16, color=GREEN, bold=True)
    add_textbox(slide, x + Inches(0.25), y + Inches(0.52), Inches(5.2), Inches(0.3),
                desc, font_size=12, color=WHITE)
    add_textbox(slide, x + Inches(0.25), y + Inches(0.82), Inches(5.2), Inches(0.25),
                detail, font_size=11, color=GRAY)

# ═══════════════════════════════════
# Slide 7: 创新价值
# ═══════════════════════════════════
slide = prs.slides.add_slide(prs.slide_layouts[6])
add_bg(slide)
add_accent_line(slide, Inches(0.8), Inches(0.7), Inches(0.6))
add_textbox(slide, Inches(0.8), Inches(0.8), Inches(10), Inches(0.5),
            "05 / 创新价值", font_size=14, color=GREEN, bold=True)
add_textbox(slide, Inches(0.8), Inches(1.1), Inches(11), Inches(0.8),
            "为什么选择这个方向？", font_size=28, color=WHITE, bold=True)

values = [
    ("选题差异化", "体育 + AIGC 在历届大赛中占比不足 3%，草根体育历史档案活化是蓝海方向。"),
    ("情感共鸣强", "真实人物、真实故事、十年队史，让观众在互动中感受体育精神与青春记忆。"),
    ("技术融合深", "文本、图像、交互、数据可视化多模态 AIGC 深度融合，形成完整产品形态。"),
    ("社会价值高", "记录和传承民间体育文化，响应体育强国、全民健身战略。"),
    ("可扩展性强", "叙事数据结构通用化，可快速适配其他球队、赛事、体育项目。"),
]

for i, (title, desc) in enumerate(values):
    y = Inches(2.0) + Inches(0.95) * i
    add_rect(slide, Inches(0.8), y, Inches(11.7), Inches(0.8), BG_CARD, radius=True)
    add_textbox(slide, Inches(1.0), y + Inches(0.12), Inches(2.5), Inches(0.3),
                title, font_size=15, color=GREEN, bold=True)
    add_textbox(slide, Inches(3.5), y + Inches(0.12), Inches(8.7), Inches(0.5),
                desc, font_size=13, color=GRAY)

# ═══════════════════════════════════
# Slide 8: 结尾
# ═══════════════════════════════════
slide = prs.slides.add_slide(prs.slide_layouts[6])
add_bg(slide)
add_rect(slide, Inches(0), Inches(0), Inches(13.333), Inches(0.08), GREEN)
add_rect(slide, Inches(0), Inches(7.42), Inches(13.333), Inches(0.08), GREEN)

add_textbox(slide, Inches(0), Inches(2.2), Inches(13.333), Inches(1.0),
            "绿茵纪", font_size=72, color=WHITE, bold=True, align=PP_ALIGN.CENTER)
add_textbox(slide, Inches(0), Inches(3.4), Inches(13.333), Inches(0.6),
            "用 AI 记录体育历史，让草根故事被看见", font_size=24, color=GREEN, align=PP_ALIGN.CENTER)
add_textbox(slide, Inches(0), Inches(4.6), Inches(13.333), Inches(0.4),
            "感谢聆听", font_size=28, color=GRAY, align=PP_ALIGN.CENTER)
add_textbox(slide, Inches(0), Inches(5.4), Inches(13.333), Inches(0.3),
            "新闻男足，快乐就完事了！", font_size=16, color=GRAY, align=PP_ALIGN.CENTER)

# 保存
output_path = os.path.join(os.path.dirname(os.path.abspath(__file__)), '..', '绿茵纪-AIGC大赛答辩PPT.pptx')
os.makedirs(os.path.dirname(output_path), exist_ok=True)
prs.save(output_path)
print(f"PPT 已生成: {output_path}")
