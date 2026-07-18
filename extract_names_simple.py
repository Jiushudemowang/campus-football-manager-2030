
# 读取提取的文本
with open('新闻男足简史_全文.txt', 'r', encoding='utf-8') as f:
    text = f.read()

# 从文档中我们已经知道的明确提到的人物
confirmed_names = [
    # 2016级
    '吴志鸣', '陈子昱', '刘梦凡', '肖潇', '曹新宇', '马金瑞', '张豪',
    '马晨阳', '陈涛', '李建昕', '孙啸峰', '王戈', '管天浩',

    # 2017级
    '欧翔', '朱家毅', '陈然', '向一丞', '邹思贤', '文若愚',

    # 其他重要人物
    '吾尔肯·努热迪勒', '阿负', '吴杨楚涵', '杨云帆', '王俊博'
]

# 统计这些名字的出现次数
name_count = {}
for name in confirmed_names:
    count = text.count(name)
    name_count[name] = count

# 按频率排序
sorted_names = sorted(name_count.items(), key=lambda x: x[1], reverse=True)

print('《新闻男足简史》中明确提到的人物：')
print('='*50)
print(f'共 {len(confirmed_names)} 人\n')
for name, count in sorted_names:
    print(f'{name}: 出现 {count} 次')

# 也尝试找更多人物
# 读取文档中的表格
from docx import Document
doc = Document('新闻男足简史.docx')

print(f'\n文档中有 {len(doc.tables)} 个表格，让我们查看表格内容：')
print('='*50)

all_table_names = set()

for i, table in enumerate(doc.tables):
    print(f'\n表格 {i+1}:')
    for row in table.rows:
        row_text = []
        for cell in row.cells:
            cell_text = cell.text.strip()
            if cell_text:
                row_text.append(cell_text)
                # 检查单元格是否是人名（2-4字，不是常见词）
                if 2 &lt;= len(cell_text) &lt;= 4 and cell_text not in ['队长', '副队长', '成员', '后续加入']:
                    all_table_names.add(cell_text)
        if row_text:
            print(' | '.join(row_text))

print(f'\n从表格中找到的名字（可能包含队员名单）：')
print(sorted(all_table_names))

# 保存结果
with open('人物统计结果.txt', 'w', encoding='utf-8') as f:
    f.write('《新闻男足简史》人物统计\n')
    f.write('='*50 + '\n\n')
    f.write('明确提到的核心人物（按出现频率）：\n')
    for name, count in sorted_names:
        f.write(f'{name}: {count}次\n')
    f.write(f'\n共 {len(confirmed_names)} 位核心人物\n\n')
    f.write('从表格中找到的其他名字：\n')
    f.write(', '.join(sorted(all_table_names)))

print('\n结果已保存到 人物统计结果.txt')
